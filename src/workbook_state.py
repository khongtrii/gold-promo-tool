"""Serialized, atomic updates to the shared master workbook."""

from contextlib import contextmanager
from datetime import datetime
import errno
import os
from pathlib import Path
import tempfile
import time

from openpyxl import load_workbook


@contextmanager
def workbook_write_lock(path: Path, timeout: float = 5.0):
    """Lock a stable sidecar, shared by all catalogue writers.

    The sidecar remains on disk: deleting it would allow competing processes
    to lock different files. The OS releases the lock even after a crash.
    """
    path = Path(path).resolve()
    lock_path = path.with_name(f".{path.name}.write.lock")
    with lock_path.open("a+b") as lock_file:
        lock_file.seek(0, os.SEEK_END)
        if lock_file.tell() == 0:
            lock_file.write(b"0")
            lock_file.flush()
        deadline = time.monotonic() + timeout
        while True:
            lock_file.seek(0)
            try:
                if os.name == "nt":
                    import msvcrt
                    msvcrt.locking(lock_file.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except OSError as error:
                if error.errno not in {errno.EACCES, errno.EAGAIN, errno.EDEADLK}:
                    raise
                if time.monotonic() >= deadline:
                    raise TimeoutError(
                        f"Master workbook is being updated: {path}. Please try again."
                    ) from error
                time.sleep(0.05)
        try:
            yield
        finally:
            lock_file.seek(0)
            if os.name == "nt":
                msvcrt.locking(lock_file.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)


def append_sitegroup_history(path: Path, rows: set[tuple[str, str, str]]) -> int:
    """Append unique catalogue/group/site triples without rewriting history."""
    path = Path(path)
    original_headers = ("CATALOGUE", "SITEGROUP", "SITE")
    headers = (*original_headers, "CREATE_AT")
    with workbook_write_lock(path):
        workbook = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
        temporary_path = None
        try:
            created = "SITE_GROUP_CHECK" not in workbook.sheetnames
            if created:
                sheet = workbook.create_sheet("SITE_GROUP_CHECK")
                sheet.append(headers)
                migrated = False
            else:
                sheet = workbook["SITE_GROUP_CHECK"]
                actual = tuple(cell.value for cell in sheet[1])
                migrated = actual == original_headers
                if migrated:
                    sheet.cell(1, 4, "CREATE_AT")
                elif actual != headers:
                    raise ValueError(
                        "SITE_GROUP_CHECK must have exactly these columns: "
                        "CATALOGUE, SITEGROUP, SITE, CREATE_AT."
                    )
            existing = {
                tuple("" if value is None else str(value).strip() for value in row)
                for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True)
            }
            additions = sorted(rows - existing)
            if not created and not migrated and not additions:
                return 0
            created_at = datetime.now() if additions else None
            for row in additions:
                sheet.append((*row, created_at))
                for cell in sheet[sheet.max_row][:3]:
                    cell.data_type = "s"
                    cell.number_format = "@"
                sheet.cell(sheet.max_row, 4).number_format = "dd/mm/yyyy hh:mm:ss"
            with tempfile.NamedTemporaryFile(
                dir=path.parent, suffix=path.suffix, delete=False
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            workbook.close()
            os.replace(temporary_path, path)
            return len(additions)
        finally:
            workbook.close()
            if workbook.vba_archive is not None:
                workbook.vba_archive.close()
            if temporary_path is not None and temporary_path.exists():
                temporary_path.unlink()
