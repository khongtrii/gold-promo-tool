import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook

from src.constant.required import required_cm, required_stage2
from src.service.template_service import Template_ETL


class DiscountParsingTest(unittest.TestCase):
    def test_validate_splits_delete_rows_with_or_without_attribute_check(self):
        for check_attribute in (False, True):
            with self.subTest(check_attribute=check_attribute):
                data = pd.DataFrame({column: ["1"] * 5 for column in required_cm})
                data["GOLD PROMO TYPE"] = "GOLD PROMO"
                data["ATTRIBUTE MARKETING"] = ["delete hero", "please DELETE", "unknown", "", "Hero"]
                data["PURCHASE VAT"] = "10%"
                data["GOLD CODE"] = ["GC1", "GC2", "GC3", "GC4", "GC5"]
                data["SO"] = ["SO1", "SO2", "SO3", "SO4", "SO5"]
                data.loc[:1, "NORMAL PURCHASE PRICE"] = "invalid"
                etl = Template_ETL([Path("source.xlsx")], check_attribute=check_attribute)
                etl.dept = {"source.xlsx": "110"}
                with patch.object(etl, "_load_source_metadata"), patch(
                    "src.service.template_service.pd.read_excel", return_value=data
                ):
                    etl._load_src()
                self.assertEqual(etl.src["GOLD CODE"].tolist(), ["GC3", "GC4", "GC5"])
                self.assertEqual(etl.src["_SOURCE_ROW"].tolist(), [10, 11, 12])
                self.assertEqual(etl.gold_code_delete["GOLD CODE"].tolist(), ["GC1", "GC2"])
                self.assertEqual(etl.gold_code_delete["SO"].tolist(), ["SO1", "SO2"])
                notes = etl.src["NOTE ERR FROM MASTER DATA"].fillna("")
                if check_attribute:
                    self.assertIn(etl.ATTRIBUTE_MARKETING_ERROR, notes.iloc[0])
                    self.assertIn(etl.ATTRIBUTE_MARKETING_ERROR, notes.iloc[1])
                else:
                    self.assertEqual(notes.tolist(), ["", "", ""])
                    self.assertEqual(etl.src["ATTRIBUTE MARKETING"].tolist(), ["unknown", "", "Hero"])

    def test_pp_dates_follow_delivery_type_order_dates_from_plan(self):
        etl = Template_ETL([])
        etl.cata = "C01"
        etl.plan = pd.DataFrame(
            {
                "CROSS-DOCKING | ORDER DATE 1": [pd.Timestamp("2026-10-01")],
                "CROSS-DOCKING | ORDER DATE 3": [pd.Timestamp("2026-10-03")],
                "DIRECT | ORDER DATE 1": [pd.Timestamp("2026-10-04")],
                "DIRECT | ORDER DATE 3": [pd.Timestamp("2026-10-06")],
                "VINAMILK | ORDER DATE 1": [pd.Timestamp("2026-10-07")],
                "VINAMILK | ORDER DATE 3": [pd.Timestamp("2026-10-09")],
            }
        )
        data = pd.DataFrame(
            {
                "DELIVERY TYPE": [
                    "CROSS-DOCKING",
                    "DIRECT",
                    "VINAMILK",
                    "DIRECT",
                ],
                "PP START DATE": pd.to_datetime(
                    ["2026-10-01", "2026-10-05", "2026-10-07", None]
                ),
                "PP END DATE": pd.to_datetime(
                    ["2026-10-03", "2026-10-05", "2026-10-10", None]
                ),
            }
        )

        result = etl._validate_pp_dates_against_plan(data)

        self.assertEqual(result.at[0, "NOTE ERR FROM MASTER DATA"], "")
        self.assertIn(
            "PP START DATE phải bằng DIRECT | ORDER DATE 1 (04/10/2026).",
            result.at[1, "NOTE ERR FROM MASTER DATA"],
        )
        self.assertIn(
            "PP END DATE phải lớn hơn hoặc bằng DIRECT | ORDER DATE 3 (06/10/2026).",
            result.at[1, "NOTE ERR FROM MASTER DATA"],
        )
        self.assertEqual(result.at[2, "NOTE ERR FROM MASTER DATA"], "")
        self.assertEqual(result.at[3, "NOTE ERR FROM MASTER DATA"], "")

    def test_blank_contract_reports_required_data_and_skips_contract_normalization(self):
        etl = Template_ETL([])
        etl.dict_network = {"wh": []}
        data = pd.DataFrame(
            {
                "COMMERCIAL CONTRACT": [None, "   ", "ABCD1234"],
                "NOTE ERR FROM MASTER DATA": ["", "", ""],
            }
        )

        etl._check_required_data(data, ["COMMERCIAL CONTRACT"])
        data["COMMERCIAL CONTRACT"] = data["COMMERCIAL CONTRACT"].map(
            etl._contract_checking
        )

        self.assertEqual(data["COMMERCIAL CONTRACT"].tolist(), ["", "", "ABCD"])
        self.assertEqual(
            data["NOTE ERR FROM MASTER DATA"].tolist(),
            [
                "Các cột bắt buộc đang để trống: COMMERCIAL CONTRACT",
                "Các cột bắt buộc đang để trống: COMMERCIAL CONTRACT",
                "",
            ],
        )

    def test_pipeline_normalizes_contract_before_duplicate_validation(self):
        etl = Template_ETL([])
        etl.dict_network = {"wh": [], "store": []}
        etl.dept = {"source.xlsx": "110"}
        etl.should_generate_so = False
        etl.src = pd.DataFrame(
            {
                "FILE NAME": ["source.xlsx", "source.xlsx"],
                "STRUCTURE": ["110", "110"],
                "SO": ["", ""],
                "GOLD CODE": ["GC1", "GC1"],
                "LV": ["1", "1"],
                "LU": ["1", "1"],
                "SUPPLIER CODE": ["SUP1", "SUP1"],
                "COMMERCIAL CONTRACT": ["R038", "R0380"],
                "PURCHASE NETWORK": ["901", "901"],
                "NORMAL PURCHASE PRICE": ["100", "100"],
                "DISCOUNT (% OR VALUE)": ["10%", "10%"],
            }
        )

        passthrough_methods = (
            "_check_network",
            "_ppNetwork_gpNetwork",
            "_field_validator",
            "_validate_structure_gold_lv",
            "_check_allocation",
            "_convert_date",
            "_validate_pp_dates_against_plan",
            "_validate_overlapping_price_or_discount",
        )
        patchers = [
            patch.object(etl, name, side_effect=lambda data: data)
            for name in passthrough_methods
        ]
        duplicate_result = {}

        def validate_duplicates_then_stop(data):
            duplicate_result["data"] = Template_ETL._validate_duplicate_purchase_information(
                etl, data
            )
            raise RuntimeError("stop after duplicate validation")

        patchers.append(
            patch.object(
                etl,
                "_validate_duplicate_purchase_information",
                side_effect=validate_duplicates_then_stop,
            )
        )
        for patcher in patchers:
            patcher.start()
        try:
            with self.assertRaisesRegex(RuntimeError, "stop after duplicate validation"):
                etl._pipeline()
        finally:
            for patcher in reversed(patchers):
                patcher.stop()

        result = duplicate_result["data"]
        self.assertEqual(result["COMMERCIAL CONTRACT"].tolist(), ["R038", "R038"])
        self.assertTrue(
            result["NOTE ERR FROM MASTER DATA"].str.contains(
                "Thông tin mua hàng và chiết khấu bị trùng", regex=False
            ).all()
        )

    def test_discount_difference_allows_only_type_1_or_2_with_type_3(self):
        allowed = [
            pd.Series(["10%", "10+2"]),
            pd.Series(["10000", "10TH+2TH"]),
        ]
        rejected = [
            pd.Series(["10%", "20%"]),
            pd.Series(["10000", "20000"]),
            pd.Series(["10+2", "20TH+3TH"]),
            pd.Series(["10%", "10000"]),
            pd.Series(["10%", "20%", "10+2"]),
        ]

        for discounts in allowed:
            with self.subTest(discounts=discounts.tolist()):
                self.assertFalse(Template_ETL._has_invalid_discount_difference(discounts))
        for discounts in rejected:
            with self.subTest(discounts=discounts.tolist()):
                self.assertTrue(Template_ETL._has_invalid_discount_difference(discounts))

    def test_overlapping_network_applies_discount_type_exception(self):
        etl = Template_ETL([])
        base = {
            "GOLD CODE": ["02043862", "02043862"],
            "LV": ["1", "1"],
            "PURCHASE NETWORK EXPANDED": ["1001", "1001"],
            "NORMAL PURCHASE PRICE": ["10000", "10000"],
        }

        allowed = etl._validate_overlapping_price_or_discount(
            pd.DataFrame(base | {"DISCOUNT (% OR VALUE)": ["10%", "10+2"]})
        )
        rejected = etl._validate_overlapping_price_or_discount(
            pd.DataFrame(base | {"DISCOUNT (% OR VALUE)": ["10%", "10000"]})
        )

        self.assertTrue(allowed["NOTE ERR FROM MASTER DATA"].eq("").all())
        self.assertTrue(
            rejected["NOTE ERR FROM MASTER DATA"].str.contains(
                "Thông tin mua hàng và chiết khấu bị trùng", regex=False
            ).all()
        )

    def test_missing_allocations_are_combined_for_every_row_in_group(self):
        etl = Template_ETL([])
        etl.dict_network = {"store": ["101", "102"]}
        data = pd.DataFrame(
            {
                "GOLD CODE": ["GC1", "GC1"],
                "LV": ["1", "1"],
                "LU": ["1", "1"],
                "SUPPLIER CODE": ["SUP1", "SUP2"],
                "COMMERCIAL CONTRACT": ["CON1", "CON2"],
                "PURCHASE NETWORK EXPANDED": ["101", "102"],
                "% DELIVERY 1": ["100", "100"],
                "% DELIVERY 2": ["", ""],
                "% DELIVERY 3": ["", ""],
                # Values on another row must not satisfy the current row's
                # PURCHASE NETWORK allocation requirement.
                "101": ["", "1"],
                "102": ["1", ""],
            }
        )

        result = etl._check_allocation(data)

        expected = (
            "Thiếu phân bổ đối với các cửa hàng: 101;102 "
            "dựa trên PURCHASE NETWORK EXPANDED."
        )
        self.assertEqual(result["NOTE ERR FROM MASTER DATA"].tolist(), [expected, expected])

    def test_blank_discount_defaults_to_zero_percent(self):
        data = pd.DataFrame({"DISCOUNT (% OR VALUE)": [None, "", "  ", "10%"]})

        result = Template_ETL._default_blank_discounts(data)

        self.assertEqual(
            result["DISCOUNT (% OR VALUE)"].tolist(),
            ["0%", "0%", "0%", "10%"],
        )

    def test_accepts_supported_discount_formats(self):
        valid_values = [
            "10",
            "10.5",
            "10,5",
            "10%",
            "10 + 2",
            "10T + 2T",
            "10th+2TH",
            "10t+2th",
        ]

        for value in valid_values:
            normalized = value.replace(" ", "")
            with self.subTest(value=value):
                self.assertIsNotNone(
                    Template_ETL.VALID_DISCOUNT.fullmatch(normalized)
                )

    def test_rejects_percentage_free_goods_and_partial_unit_formats(self):
        invalid_values = [
            "10%+2%",
            "10%+2",
            "10+2%",
            "10T+2",
            "10+2TH",
            "abc",
            "",
        ]

        for value in invalid_values:
            with self.subTest(value=value):
                self.assertIsNone(Template_ETL.VALID_DISCOUNT.fullmatch(value))

    def test_only_plain_number_plus_number_is_non_warehouse(self):
        data = pd.DataFrame(
            {
                "GOLD CODE": [
                    "NORMAL",
                    "NORMAL",
                    "NORMAL",
                    "NORMAL",
                    "NORMAL",
                    "NORMAL",
                    "NORMAL",
                ],
                "DISCOUNT (% OR VALUE)": [
                    "10+2",
                    "10 + 2",
                    "10T+2T",
                    "10TH+2TH",
                    "10%+2%",
                    "10%",
                    "10",
                ]
            }
        )

        result = Template_ETL._get_non_warehouse_src(data)

        self.assertEqual(result.index.tolist(), [0, 1])

    def test_warehouse_exception_is_not_reported_as_non_warehouse(self):
        data = pd.DataFrame(
            {
                "GOLD CODE": ["02043862", "99999999"],
                "DISCOUNT (% OR VALUE)": ["10+2", "10+2"],
            }
        )

        result = Template_ETL._get_non_warehouse_src(data, {"02043862"})

        self.assertEqual(result["GOLD CODE"].tolist(), ["99999999"])

    def test_exception_gold_codes_reject_all_carton_discount_variants(self):
        values = [
            "10T+2T",
            "10TH+2TH",
            "10T+2TH",
            "10th + 2t",
            "10+2",
            "10",
            "10%",
            "10TH+2TH",
        ]
        etl = Template_ETL([], exception_discount_gold_codes={"02043862", "CUSTOM01"})
        etl.src = pd.DataFrame(
            {
                "GOLD CODE": ["02043862"] * 7 + ["99999999"],
                "DISCOUNT (% OR VALUE)": values,
                "NOTE ERR FROM MASTER DATA": [""] * len(values),
            }
        )

        etl._validate_exception_carton_discounts()

        notes = etl.src["NOTE ERR FROM MASTER DATA"]
        self.assertTrue(notes.iloc[:4].eq(etl.EXCEPTION_CARTON_DISCOUNT_ERROR).all())
        self.assertTrue(notes.iloc[4:].eq("").all())

    def test_exception_carton_validation_tracks_the_current_exception_list(self):
        etl = Template_ETL([], exception_discount_gold_codes={"CUSTOM01"})
        etl.src = pd.DataFrame(
            {
                "GOLD CODE": ["CUSTOM01"],
                "DISCOUNT (% OR VALUE)": ["10TH+2TH"],
                "NOTE ERR FROM MASTER DATA": ["Existing error"],
            }
        )

        etl._validate_exception_carton_discounts()
        self.assertEqual(
            etl.src.at[0, "NOTE ERR FROM MASTER DATA"],
            "Existing error | " + etl.EXCEPTION_CARTON_DISCOUNT_ERROR,
        )

        etl.exception_discount_gold_codes = set()
        etl._validate_exception_carton_discounts()
        self.assertEqual(etl.src.at[0, "NOTE ERR FROM MASTER DATA"], "Existing error")

    def test_invalid_site_rule_skips_following_network_checks_for_its_group(self):
        etl = Template_ETL([])
        etl.dict_network = {"network": {}, "store": ["1001"]}
        data = pd.DataFrame(
            {
                "FILE NAME": ["source.xlsx"],
                "GOLD CODE": ["02043862"],
                "LV": ["1"],
                "LU": ["1"],
                "PURCHASE NETWORK": ["1001++"],
                "GOLD PROMO NETWORK": ["9999"],
            }
        )

        result = etl._check_network(data)
        result = etl._ppNetwork_gpNetwork(result)
        note = result.at[0, "NOTE ERR FROM MASTER DATA"]

        self.assertIn("PURCHASE NETWORK không đúng SITE RULE.", note)
        self.assertNotIn("Cửa hàng", note)
        self.assertNotIn("Thiếu cửa hàng", note)
        self.assertNotIn("Dư cửa hàng", note)

    def test_subtracted_literal_sites_must_belong_to_the_positive_network(self):
        etl = Template_ETL([])
        etl.dict_network = {
            "network": {"8200": "101;102;103", "8210": "101;102"},
            "store_hyper": ["101", "102", "103"],
            "store": ["101", "102", "103"],
        }
        data = pd.DataFrame(
            {
                "FILE NAME": ["source.xlsx", "source.xlsx"],
                "GOLD CODE": ["GC1", "GC2"],
                "LV": ["1", "1"],
                "LU": ["1", "1"],
                "PURCHASE NETWORK": ["8200(-101;148,806)", "8210-101"],
                "GOLD PROMO NETWORK": ["8200-101", "8210-153"],
            }
        )

        result = etl._check_network(data)

        self.assertIn(
            "PURCHASE NETWORK có site trừ không thuộc network tổng: 148;806.",
            result.at[0, "NOTE ERR FROM MASTER DATA"],
        )
        self.assertNotIn(
            "PURCHASE NETWORK có site trừ",
            result.at[1, "NOTE ERR FROM MASTER DATA"],
        )
        self.assertIn(
            "GOLD PROMO NETWORK có site trừ không thuộc network tổng: 153.",
            result.at[1, "NOTE ERR FROM MASTER DATA"],
        )

    def test_so_longer_than_ten_characters_is_reported(self):
        etl = Template_ETL([])
        data = pd.DataFrame(
            {
                "SO": ["C01D410-01", "C001D410-01", "  C01D41-01  ", None],
                "NOTE ERR FROM MASTER DATA": ["", "Existing error", "", ""],
            }
        )

        result = etl._validate_so_length(data)

        self.assertEqual(result.at[0, "NOTE ERR FROM MASTER DATA"], "")
        self.assertEqual(
            result.at[1, "NOTE ERR FROM MASTER DATA"],
            "Existing error | SO chỉ được tối đa 10 ký tự.",
        )
        self.assertEqual(result.at[2, "NOTE ERR FROM MASTER DATA"], "")
        self.assertEqual(result.at[3, "NOTE ERR FROM MASTER DATA"], "")

    def test_network_8000_skips_all_network_validation_for_its_group(self):
        etl = Template_ETL([])
        etl.dict_network = {"network": {}, "store": ["1001"]}
        data = pd.DataFrame(
            {
                "FILE NAME": ["source.xlsx", "source.xlsx"],
                "GOLD CODE": ["02043862", "02043863"],
                "LV": ["1", "1"],
                "LU": ["1", "1"],
                "PURCHASE NETWORK": ["8000", "1001"],
                "GOLD PROMO NETWORK": ["9999", "8000"],
            }
        )

        result = etl._check_network(data)
        result = etl._ppNetwork_gpNetwork(result)
        for index in result.index:
            self.assertEqual(result.at[index, "NOTE ERR FROM MASTER DATA"], "Không được sử dụng network 8000.")
            self.assertEqual(result.at[index, "PURCHASE NETWORK EXPANDED"], "")
            self.assertEqual(result.at[index, "GOLD PROMO NETWORK EXPANDED"], "")

    def test_restores_percentage_formatted_source_cells(self):
        columns = [
            "DISCOUNT (% OR VALUE)",
            "% DELIVERY 1",
            "% DELIVERY 2",
            "% DELIVERY 3",
        ]
        with TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Template"
            for column_index, column in enumerate(columns, start=1):
                sheet.cell(7, column_index, column)
            for column_index, value in enumerate([0.5, 0.5, 0.3, 0.2], start=1):
                cell = sheet.cell(8, column_index, value)
                cell.number_format = "0%"
            sheet.cell(9, 1, "50%")
            workbook.save(path)

            data = pd.DataFrame(
                {
                    "DISCOUNT (% OR VALUE)": ["0.5", "50%"],
                    "% DELIVERY 1": ["0.5", ""],
                    "% DELIVERY 2": ["0.3", ""],
                    "% DELIVERY 3": ["0.2", ""],
                    "NORMAL DISCOUNT ON PURCHASE (%)": ["0.5", "0.5"],
                }
            )
            Template_ETL._restore_percentage_cells(data, path)

        self.assertEqual(data.loc[0, "DISCOUNT (% OR VALUE)"], "50%")
        self.assertEqual(data.loc[0, "% DELIVERY 1"], "50%")
        self.assertEqual(data.loc[0, "% DELIVERY 2"], "30%")
        self.assertEqual(data.loc[0, "% DELIVERY 3"], "20%")
        self.assertEqual(data.loc[1, "DISCOUNT (% OR VALUE)"], "50%")
        self.assertEqual(data.loc[0, "NORMAL DISCOUNT ON PURCHASE (%)"], "0.5")
        self.assertIsNotNone(Template_ETL.VALID_DISCOUNT.fullmatch(data.loc[0, "DISCOUNT (% OR VALUE)"]))

        allocation = data.iloc[[0]].copy()
        allocation["GOLD CODE"] = "02043862"
        allocation["LV"] = "1"
        allocation["LU"] = "1"
        allocation["SUPPLIER CODE"] = "SUP"
        allocation["COMMERCIAL CONTRACT"] = "CONT"
        allocation["PURCHASE NETWORK EXPANDED"] = ""
        result = Template_ETL([])._check_allocation(allocation)
        self.assertNotIn("Tổng % DELIVERY", result.at[0, "NOTE ERR FROM MASTER DATA"])

    def test_stage2_validates_vat_and_normalizes_integral_sale_price(self):
        data = pd.DataFrame(
            {
                "SALE VAT": ["10%", " kkkt ", "invalid", "0.1"],
                "PROMOTION SALE PRICE": [
                    "45100.0",
                    "45100.0000000001",
                    "45100.5",
                    "not a number",
                ],
            }
        )

        result = Template_ETL([])._validate_stage2_sale_values(data)

        self.assertEqual(result["SALE VAT"].tolist()[:2], ["10%", "KKKT"])
        self.assertEqual(result["PROMOTION SALE PRICE"].tolist()[:2], ["45100", "45100"])
        self.assertIn("SALE VAT chỉ được phép", result.at[2, "NOTE ERR FROM MASTER DATA"])
        self.assertNotIn("SALE VAT chỉ được phép", result.at[3, "NOTE ERR FROM MASTER DATA"])
        self.assertEqual(result.at[3, "SALE VAT"], "10%")
        self.assertIn("PROMOTION SALE PRICE phải là số nguyên", result.at[2, "NOTE ERR FROM MASTER DATA"])
        self.assertIn("PROMOTION SALE PRICE phải là số nguyên", result.at[3, "NOTE ERR FROM MASTER DATA"])

    def test_gold_promo_type_filters_out_of_promo_and_retains_invalid_rows(self):
        data = pd.DataFrame(
            {
                "GOLD PROMO TYPE": [
                    "GOLD PROMO",
                    "OUT OF GOLD PROMO",
                    " GOLD PROMO ",
                    "gold promo",
                    "INVALID",
                    "",
                ]
            }
        )
        etl = Template_ETL([])

        result = etl._filter_gold_promo_rows(data)

        self.assertEqual(result.index.tolist(), [0, 2, 3, 4, 5])
        self.assertEqual(result.loc[[0, 2], "GOLD PROMO TYPE"].tolist(), ["GOLD PROMO"] * 2)
        self.assertTrue(etl.has_invalid_gold_promo_type)
        for index in (3, 4, 5):
            self.assertIn(
                "GOLD PROMO TYPE chỉ được phép",
                result.at[index, "NOTE ERR FROM MASTER DATA"],
            )

    def test_stage2_sp_dates_come_from_source_and_match_plan_dates(self):
        data = pd.DataFrame({column: ["1", "1", "1"] for column in required_stage2})
        data["GOLD PROMO TYPE"] = "GOLD PROMO"
        data["PRICELIST"] = ["20300001", "20100001", "20300002"]
        data["SP START DAY"] = ["2", "3", "4"]
        data["SP START MONTH"] = "10"
        data["SP START YEAR"] = "2099"
        data["SP END DAY"] = ["10", "11", "12"]
        data["SP END MONTH"] = "10"
        data["SP END YEAR"] = "2099"
        data["NOTE ERR FROM MASTER DATA"] = ""
        data["FILE NAME"] = "source.xlsx"
        data["_SOURCE_ROW"] = [8, 9, 10]
        data["STRUCTURE"] = "110"
        etl = Template_ETL([])
        etl.cata = "C01"
        etl.src_listoff = data
        etl.plan = pd.DataFrame(
            {
                "CATALOGUE START DATE": [pd.Timestamp("2099-10-02")],
                "CATALOGUE END DATE": [pd.Timestamp("2099-10-10")],
            }
        )

        etl._pipeline2()

        self.assertEqual(etl.src_listoff.at[0, "SP START DATE"], pd.Timestamp("2099-10-02"))
        self.assertEqual(etl.src_listoff.at[0, "SP END DATE"], pd.Timestamp("2099-10-10"))
        self.assertEqual(etl.src_listoff.at[0, "NOTE ERR FROM MASTER DATA"], "")
        self.assertNotIn("CATALOGUE START DATE", etl.src_listoff.at[1, "NOTE ERR FROM MASTER DATA"])
        self.assertNotIn("CATALOGUE END DATE", etl.src_listoff.at[1, "NOTE ERR FROM MASTER DATA"])
        self.assertIn("CATALOGUE START DATE", etl.src_listoff.at[2, "NOTE ERR FROM MASTER DATA"])
        self.assertIn("CATALOGUE END DATE", etl.src_listoff.at[2, "NOTE ERR FROM MASTER DATA"])

    def test_stage2_loads_sp_period_from_source_instead_of_catalogue_metadata(self):
        data = pd.DataFrame({column: ["1"] for column in required_stage2})
        data["GOLD PROMO TYPE"] = "GOLD PROMO"
        data["SALE VAT"] = "10%"
        data["PROMOTION SALE PRICE"] = "100"
        data["PRICELIST"] = "20100001"
        data["SP START DAY"] = "2"
        data["SP START MONTH"] = "10"
        data["SP START YEAR"] = "2099"
        data["SP END DAY"] = "10"
        data["SP END MONTH"] = "10"
        data["SP END YEAR"] = "2099"
        etl = Template_ETL([Path("source.xlsx")])
        etl.dept = {"source.xlsx": "110"}

        with patch.object(etl, "_load_source_metadata", return_value={}), patch(
            "src.service.template_service.pd.read_excel", return_value=data
        ):
            etl._load_src_listoff()

        self.assertEqual(
            etl.src_listoff.loc[0, list(etl.SP_PERIOD_COLUMNS)].tolist(),
            ["2", "10", "2099", "10", "10", "2099"],
        )

    def test_validate_requires_valid_purchase_vat_with_or_without_attribute_check(self):
        invalid_values = ["", "   ", None, float("nan"), "7%", "10", "invalid", "0.07", "inf", "NaN"]
        valid_values = [
            "0%", "5%", "8%", "10%", " kkkt ", "kct",
            0, 0.05, 0.08, 0.1, "0.0", "0.050", " 0.08 ", "0.10", "0,08",
        ]
        expected = [
            "0%", "5%", "8%", "10%", "KKKT", "KCT",
            "0%", "5%", "8%", "10%", "0%", "5%", "8%", "10%", "8%",
        ]
        values = invalid_values + valid_values
        sale = pd.DataFrame({"SALE VAT": values, "PROMOTION SALE PRICE": "100"})
        Template_ETL([])._validate_stage2_sale_values(sale)
        sale_notes = sale["NOTE ERR FROM MASTER DATA"].fillna("").tolist()
        for note in sale_notes[:len(invalid_values)]:
            self.assertIn("SALE VAT chỉ được phép", note)
        self.assertEqual(sale_notes[len(invalid_values):], [""] * len(valid_values))
        self.assertEqual(sale["SALE VAT"].tolist()[len(invalid_values):], expected)
        for check_attribute in (False, True):
            with self.subTest(check_attribute=check_attribute):
                data = pd.DataFrame({column: ["1"] * len(values) for column in required_cm})
                data["GOLD PROMO TYPE"] = "GOLD PROMO"
                data["ATTRIBUTE MARKETING"] = "Hero"
                data["PURCHASE VAT"] = values
                etl = Template_ETL([Path("source.xlsx")], check_attribute=check_attribute)
                etl.dept = {"source.xlsx": "110"}
                with patch.object(etl, "_load_source_metadata"), patch(
                    "src.service.template_service.pd.read_excel", return_value=data
                ):
                    etl._load_src()
                notes = etl.src["NOTE ERR FROM MASTER DATA"].fillna("").tolist()
                for note in notes[:len(invalid_values)]:
                    self.assertIn(
                        "PURCHASE VAT chỉ được phép là 0%, 5%, 8%, 10%, KKKT hoặc KCT.",
                        note,
                    )
                self.assertEqual(notes[len(invalid_values):], [""] * len(valid_values))
                self.assertEqual(
                    etl.src["PURCHASE VAT"].tolist()[len(invalid_values):],
                    expected,
                )

    def test_validate_allows_blank_free_product_and_converts_hero(self):
        labels = ["Front page", "Back page", "unbeat"]
        for check_attribute in (False, True):
            with self.subTest(check_attribute=check_attribute):
                data = pd.DataFrame({column: ["1"] * 3 for column in required_cm})
                data["GOLD PROMO TYPE"] = "GOLD PROMO"
                data["ATTRIBUTE MARKETING"] = labels
                data["PURCHASE VAT"] = "10%"
                data["FREE PRODUCT"] = ["", None, float("nan")]
                data["DISCOUNT (% OR VALUE)"] = ["", None, float("nan")]
                etl = Template_ETL([Path("source.xlsx")], check_attribute=check_attribute)
                etl.dept = {"source.xlsx": "110"}
                with patch.object(etl, "_load_source_metadata"), patch(
                    "src.service.template_service.pd.read_excel", return_value=data
                ):
                    etl._load_src()
                self.assertEqual(
                    etl.src["NOTE ERR FROM MASTER DATA"].fillna("").tolist(),
                    ["", "", ""],
                )
                self.assertEqual(
                    etl.src["ATTRIBUTE MARKETING"].tolist(),
                    ["HERO"] * 3 if check_attribute else labels,
                )

    def test_attribute_requires_free_product_only_when_source_discount_is_present(self):
        discounts = ["10%", "1000", "1+1", "0", "0%", "", None, float("nan"), "  ", "10%"]
        for check_attribute in (False, True):
            with self.subTest(check_attribute=check_attribute):
                data = pd.DataFrame({column: ["1"] * len(discounts) for column in required_cm})
                data["GOLD PROMO TYPE"] = "GOLD PROMO"
                data["ATTRIBUTE MARKETING"] = "Hero"
                data["PURCHASE VAT"] = "10%"
                data["FREE PRODUCT"] = ["", None, float("nan"), "  ", "", "", None, "", "  ", "Gift"]
                data["DISCOUNT (% OR VALUE)"] = discounts
                etl = Template_ETL([Path("source.xlsx")], check_attribute=check_attribute)
                etl.dept = {"source.xlsx": "110"}
                with patch.object(etl, "_load_source_metadata"), patch(
                    "src.service.template_service.pd.read_excel", return_value=data
                ):
                    etl._load_src()
                notes = etl.src["NOTE ERR FROM MASTER DATA"].fillna("").tolist()
                error = "FREE PRODUCT không được để trống khi DISCOUNT (% OR VALUE) có dữ liệu."
                self.assertEqual(notes, [error] * 5 + [""] * 5 if check_attribute else [""] * 10)


if __name__ == "__main__":
    unittest.main()
