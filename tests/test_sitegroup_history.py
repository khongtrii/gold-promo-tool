from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.desktop_app import GoldPromoApp
from src.service.template_service import Template_ETL
from src.workbook_state import append_sitegroup_history, workbook_write_lock


class SiteGroupHistoryTest(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "master.xlsx"
        workbook = Workbook()
        workbook.active.title = "Other"
        workbook.active["A1"] = "=1+2"
        workbook.save(self.path)
        workbook.close()

    def make_etl(self, catalogue="C01"):
        etl = Template_ETL([], self.path, self.path)
        etl.cata = catalogue
        etl.src = pd.DataFrame({
            "FILE NAME": ["first.xlsx", "second.xlsx", "second.xlsx"],
            "SITE GROUP": ["00123", "00123", ""],
            "GOLD PROMO NETWORK EXPANDED": ["001;002", "002;001", "003"],
        })
        etl.apply_sitegroup_suggestions([
            {"expanded_network": "003", "suggested_code": "60000"}
        ])
        return etl

    def read_rows(self):
        workbook = load_workbook(self.path)
        try:
            self.assertEqual(workbook["Other"]["A1"].value, "=1+2")
            sheet = workbook["SITE_GROUP_CHECK"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), ("CATALOGUE", "SITEGROUP", "SITE", "CREATE_AT"))
            return [row[:3] for row in sheet.iter_rows(min_row=2, values_only=True)]
        finally:
            workbook.close()

    def test_all_sources_existing_and_new_groups_text_codes_and_repeat(self):
        etl = self.make_etl()
        self.assertEqual(etl.record_sitegroup_history(), 3)
        self.assertEqual(etl.record_sitegroup_history(), 0)
        self.assertEqual(self.read_rows(), [
            ("C01", "00123", "001"), ("C01", "00123", "002"),
            ("C01", "60000", "003"),
        ])
        workbook = load_workbook(self.path)
        try:
            for row in workbook["SITE_GROUP_CHECK"].iter_rows(min_row=2):
                for cell in row[:3]:
                    self.assertEqual(cell.data_type, "s")
                    self.assertEqual(cell.number_format, "@")
                self.assertIsInstance(row[3].value, datetime)
                self.assertEqual(row[3].number_format, "dd/mm/yyyy hh:mm:ss")
            self.assertEqual(len({row[3].value for row in workbook["SITE_GROUP_CHECK"].iter_rows(min_row=2)}), 1)
        finally:
            workbook.close()

    def test_existing_timestamps_stay_unchanged_when_new_sites_are_appended(self):
        etl = self.make_etl()
        etl.record_sitegroup_history()
        workbook = load_workbook(self.path)
        old_times = {
            tuple(cell.value for cell in row[:3]): row[3].value
            for row in workbook["SITE_GROUP_CHECK"].iter_rows(min_row=2)
        }
        workbook.close()
        etl.src = pd.DataFrame({
            "SITE GROUP": ["00123", "00123"],
            "GOLD PROMO NETWORK EXPANDED": ["001", "004;005"],
        })
        self.assertEqual(etl.record_sitegroup_history(), 2)
        workbook = load_workbook(self.path)
        try:
            rows = {
                tuple(cell.value for cell in row[:3]): row[3].value
                for row in workbook["SITE_GROUP_CHECK"].iter_rows(min_row=2)
            }
            for key, old_time in old_times.items():
                self.assertEqual(rows[key], old_time)
            self.assertIsInstance(rows[("C01", "00123", "004")], datetime)
            self.assertEqual(rows[("C01", "00123", "004")], rows[("C01", "00123", "005")])
        finally:
            workbook.close()

    def test_existing_three_column_sheet_is_migrated_without_guessing_old_times(self):
        workbook = load_workbook(self.path)
        sheet = workbook.create_sheet("SITE_GROUP_CHECK")
        sheet.append(["CATALOGUE", "SITEGROUP", "SITE"])
        sheet.append(["C01", "00123", "001"])
        workbook.save(self.path)
        workbook.close()
        self.assertEqual(self.make_etl().record_sitegroup_history(), 2)
        workbook = load_workbook(self.path)
        try:
            rows = list(workbook["SITE_GROUP_CHECK"].values)
            self.assertEqual(rows[0], ("CATALOGUE", "SITEGROUP", "SITE", "CREATE_AT"))
            self.assertEqual(rows[1], ("C01", "00123", "001", None))
            self.assertIsInstance(rows[2][3], datetime)
        finally:
            workbook.close()

    def test_preserves_old_sites_and_other_catalogues(self):
        etl = self.make_etl()
        etl.record_sitegroup_history()
        etl.src = pd.DataFrame({"SITE GROUP": ["00123"], "GOLD PROMO NETWORK EXPANDED": ["004"]})
        self.assertEqual(etl.record_sitegroup_history(), 1)
        etl.cata = "C02"
        self.assertEqual(etl.record_sitegroup_history(), 1)
        rows = self.read_rows()
        self.assertEqual(len(rows), 5)
        self.assertIn(("C01", "00123", "001"), rows)
        self.assertIn(("C01", "00123", "004"), rows)
        self.assertIn(("C02", "00123", "004"), rows)

    def test_bad_headers_leave_workbook_unchanged(self):
        workbook = load_workbook(self.path)
        workbook.create_sheet("SITE_GROUP_CHECK").append(["CATALOGUE", "SITE_GROUP", "SITE"])
        workbook.save(self.path)
        workbook.close()
        original = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, "exactly these columns"):
            self.make_etl().record_sitegroup_history()
        self.assertEqual(self.path.read_bytes(), original)

    def test_marketing_does_not_write(self):
        etl = self.make_etl()
        etl.check_attribute = True
        original = self.path.read_bytes()
        self.assertEqual(etl.record_sitegroup_history(), 0)
        self.assertEqual(self.path.read_bytes(), original)

    def test_incomplete_source_is_rejected_without_partial_history(self):
        etl = self.make_etl()
        for column in ("SITE GROUP", "GOLD PROMO NETWORK EXPANDED"):
            with self.subTest(column=column):
                etl = self.make_etl()
                etl.src.loc[2, column] = ""
                original = self.path.read_bytes()
                with self.assertRaisesRegex(ValueError, "every source row"):
                    etl.record_sitegroup_history()
                self.assertEqual(self.path.read_bytes(), original)

    def test_replace_failure_preserves_original_and_retry_is_idempotent(self):
        original = self.path.read_bytes()
        etl = self.make_etl()
        with patch("src.workbook_state.os.replace", side_effect=PermissionError("File open in Excel")):
            with self.assertRaises(PermissionError):
                etl.record_sitegroup_history()
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.path.parent.glob("*.xlsx")), [self.path])
        self.assertEqual(etl.record_sitegroup_history(), 3)
        self.assertEqual(etl.record_sitegroup_history(), 0)

    def test_workbook_lock_blocks_both_writers_and_releases_on_failure(self):
        with workbook_write_lock(self.path):
            with self.assertRaises(TimeoutError):
                with workbook_write_lock(self.path, timeout=0):
                    self.fail("Competing writer acquired the lock")
        with self.assertRaisesRegex(RuntimeError, "test failure"):
            with workbook_write_lock(self.path):
                raise RuntimeError("test failure")
        with workbook_write_lock(self.path, timeout=0):
            pass

    def test_parallel_catalogues_do_not_overwrite_each_other(self):
        batches = [{("C01", "00123", "001")}, {("C02", "60000", "002")}]
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda rows: append_sitegroup_history(self.path, rows), batches))
        self.assertEqual(results, [1, 1])
        self.assertEqual(set(self.read_rows()), batches[0] | batches[1])

    def test_macro_archive_and_other_sheets_are_preserved(self):
        macro_path = self.path.with_suffix(".xlsm")
        self.path.rename(macro_path)
        macro_bytes = b"macro preservation test payload"
        with ZipFile(macro_path, "a") as archive:
            archive.writestr("xl/vbaProject.bin", macro_bytes)
        etl = self.make_etl()
        etl.path_plan = macro_path
        etl.record_sitegroup_history()
        with ZipFile(macro_path) as archive:
            self.assertEqual(archive.read("xl/vbaProject.bin"), macro_bytes)
        workbook = load_workbook(macro_path, keep_vba=True)
        try:
            self.assertEqual(workbook["Other"]["A1"].value, "=1+2")
            self.assertEqual(workbook["SITE_GROUP_CHECK"].max_row, 4)
        finally:
            workbook.close()
            workbook.vba_archive.close()

    def test_sitegroup_master_updates_use_the_same_workbook_lock(self):
        etl = self.make_etl()
        with patch("src.service.template_service.workbook_write_lock") as lock, patch.object(
            etl, "_update_sitegroup_file_locked", return_value=True
        ) as update:
            self.assertTrue(etl.update_sitegroup_file([]))
        lock.assert_called_once_with(self.path)
        update.assert_called_once_with([])

    def test_concurrent_master_update_preserves_history_and_group_sheets(self):
        workbook = load_workbook(self.path)
        workbook.create_sheet("site-group").append(["SITE_GROUP", "SITE"])
        workbook.create_sheet("site-group-activated").append(["SITE_GROUP"])
        workbook.save(self.path)
        workbook.close()
        etl = self.make_etl()
        suggestions = [{"suggested_code": "60001", "expanded_network": "004;005"}]
        with ThreadPoolExecutor(max_workers=2) as executor:
            history = executor.submit(etl.record_sitegroup_history)
            update = executor.submit(etl.update_sitegroup_file, suggestions)
            self.assertEqual(history.result(), 3)
            self.assertTrue(update.result())
        self.assertEqual(len(self.read_rows()), 3)
        workbook = load_workbook(self.path)
        try:
            self.assertEqual(list(workbook["site-group"].values), [
                ("SITE_GROUP", "SITE"), ("60001", "004"), ("60001", "005"),
            ])
            self.assertEqual(list(workbook["site-group-activated"].values), [
                ("SITE_GROUP",), ("60001",),
            ])
        finally:
            workbook.close()

    def test_json_failure_can_retry_without_duplicate_history(self):
        app = Mock()
        app.pending_etl = self.make_etl()
        app._record_used_sitegroups.side_effect = [PermissionError("JSON locked"), None]
        with patch("src.desktop_app.messagebox"):
            GoldPromoApp._complete_add_sitegroup(app)
            self.assertTrue(app.pending_etl.should_generate_so_sitegroup)
            self.assertEqual(len(self.read_rows()), 3)
            GoldPromoApp._complete_add_sitegroup(app)
        self.assertFalse(app.pending_etl.should_generate_so_sitegroup)
        self.assertEqual(len(self.read_rows()), 3)

    def test_ui_only_completes_after_history_and_json_succeed(self):
        app = Mock()
        etl = self.make_etl()
        app.pending_etl = etl
        with patch.object(etl, "record_sitegroup_history", side_effect=PermissionError("locked")), patch(
            "src.desktop_app.messagebox"
        ) as messages:
            GoldPromoApp._complete_add_sitegroup(app)
        self.assertTrue(etl.should_generate_so_sitegroup)
        app._record_used_sitegroups.assert_not_called()
        app.template_mapping_button.state.assert_called_with(["disabled"])
        app.add_sitegroup_button.state.assert_called_with(["!disabled"])
        messages.showinfo.assert_not_called()
        app._release_sitegroup_session.assert_called_once()
        with patch("src.desktop_app.messagebox") as messages:
            GoldPromoApp._complete_add_sitegroup(app)
        self.assertFalse(etl.should_generate_so_sitegroup)
        app._record_used_sitegroups.assert_called_once_with(etl)
        app.template_mapping_button.state.assert_called_with(["!disabled"])
        messages.showinfo.assert_called_once()
        self.assertEqual(len(self.read_rows()), 3)


if __name__ == "__main__":
    unittest.main()
