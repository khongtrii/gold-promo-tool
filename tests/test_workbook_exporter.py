import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.desktop_app import GoldPromoApp, WorkbookExporter


class WorkbookExporterTest(unittest.TestCase):
    def test_discount_output_numbers_continue_after_other_templates(self):
        output = Path("output")
        timestamp = "1609_1200"

        expected = {
            "template_ag": "8_TEMPLATE_AG_1609_1200.xls",
            "template_ag_raw": "8_TEMPLATE_AG_RAW_1609_1200.xls",
            "report_ag_errors": "8_REPORT_AG_ERRORS_1609_1200.xls",
            "template_dc": "9_TEMPLATE_DC_1609_1200.xls",
            "template_de": "10_TEMPLATE_DE_1609_1200.xls",
        }
        for name, file_name in expected.items():
            with self.subTest(name=name):
                self.assertEqual(
                    GoldPromoApp._output_file(output, name, timestamp),
                    output / file_name,
                )

    def test_source_error_notes_use_yellow_fill_and_red_font(self):
        with TemporaryDirectory(dir=Path.cwd()) as directory:
            source = Path(directory) / "source.xlsx"
            output = Path(directory) / "errors.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Template"
            sheet.cell(7, 1, "MULTI")
            workbook.save(source)
            workbook.close()

            WorkbookExporter.write_source_errors(
                source,
                pd.DataFrame(
                    {
                        "_SOURCE_ROW": [8],
                        "NOTE ERR FROM MASTER DATA": ["Validation error"],
                    }
                ),
                output,
            )

            result = load_workbook(output)
            try:
                result_sheet = result["Template"]
                for row in (7, 8):
                    cell = result_sheet.cell(row, 2)
                    self.assertEqual(cell.fill.fill_type, "solid")
                    self.assertEqual(cell.fill.fgColor.rgb, "00FFFF00")
                    self.assertEqual(cell.font.color.type, "rgb")
                    self.assertEqual(cell.font.color.rgb, "FFFF0000")
            finally:
                result.close()

    def test_source_shaped_exports_remove_rows_excluded_by_promo_type(self):
        with TemporaryDirectory(dir=Path.cwd()) as directory:
            source = Path(directory) / "source.xlsx"
            error_output = Path(directory) / "errors.xlsx"
            processed_output = Path(directory) / "processed.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Template"
            sheet.cell(7, 1, "GOLD CODE")
            sheet.cell(7, 2, "MULTI")
            sheet.cell(8, 1, "  GOLD-1  ")
            sheet.cell(9, 1, "OUT-1")
            sheet.cell(10, 1, "GOLD-2")
            workbook.save(source)
            workbook.close()

            kept = pd.DataFrame(
                {
                    "_SOURCE_ROW": [8, 10],
                    "NOTE ERR FROM MASTER DATA": ["Validation error", ""],
                    "STRUCTURE": ["110", "110"],
                    "SO": ["SO1", "SO2"],
                    "SITE GROUP": ["10001", "10002"],
                }
            )
            WorkbookExporter.write_source_errors(
                source,
                kept.iloc[[0]],
                error_output,
                included_source_rows=kept["_SOURCE_ROW"],
            )
            WorkbookExporter.write_processed_source(source, kept, processed_output)

            for path in (error_output, processed_output):
                result = load_workbook(path)
                try:
                    result_sheet = result["Template"]
                    self.assertEqual(result_sheet.cell(8, 1).value, "GOLD-1")
                    self.assertEqual(result_sheet.cell(9, 1).value, "GOLD-2")
                    self.assertIsNone(result_sheet.cell(10, 1).value)
                finally:
                    result.close()

    def test_combined_processed_source_excludes_internal_columns(self):
        with TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory) / "ALL_PROCESSED.xlsx"
            processed = pd.DataFrame(
                {
                    "NO": [1, 2],
                    "STRUCTURE": [110, 120],
                    "GOLD CODE": ["  A ", "B  "],
                    "NOTE ERR FROM MASTER DATA": ["", ""],
                    "FILE NAME": [" a.xlsx ", "b.xlsx  "],
                    "_SOURCE_ROW": [8, 8],
                    "SITE GROUP": ["10001", "10002"],
                    "SO": ["SO1", "SO2"],
                    "PURCHASE NETWORK EXPANDED": ["1;2", "3;4"],
                    "GOLD PROMO NETWORK EXPANDED": ["1;2", "3;4"],
                    "DISCOUNT_NETWORK_EXPANDED": ["1;2", "3;4"],
                    "PP START DATE": ["01/01/2026", "01/01/2026"],
                    "PP END DATE": ["31/01/2026", "31/01/2026"],
                }
            )

            WorkbookExporter.write_combined_processed_source(processed, output)

            result = pd.read_excel(output, sheet_name="Processed", dtype=str)
            self.assertEqual(
                result.columns.tolist(),
                ["NO", "STRUCTURE", "GOLD CODE", "SITE GROUP", "SO", "FILE NAME"],
            )
            self.assertEqual(result["FILE NAME"].tolist(), ["a.xlsx", "b.xlsx"])
            self.assertEqual(result["GOLD CODE"].tolist(), ["A", "B"])
            self.assertEqual(result["SO"].tolist(), ["SO1", "SO2"])
            self.assertEqual(len(result), 2)

    def test_excel_value_strips_only_surrounding_whitespace(self):
        self.assertEqual(WorkbookExporter._excel_value("  A  B  "), "A  B")

    def test_combined_processed_source_omits_sitegroup_before_add_sitegroup(self):
        with TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory) / "ALL_PROCESSED.xlsx"
            processed = pd.DataFrame(
                {
                    "STRUCTURE": [110],
                    "SITE GROUP": ["10001"],
                    "SO": ["SO1"],
                    "FILE NAME": ["a.xlsx"],
                }
            )

            WorkbookExporter.write_combined_processed_source(
                processed,
                output,
                include_sitegroup=False,
            )

            result = pd.read_excel(output, sheet_name="Processed", dtype=str)
            self.assertEqual(result.columns.tolist(), ["STRUCTURE", "SO", "FILE NAME"])

    def test_export_processed_source_adds_combined_file_for_multiple_sources(self):
        app = GoldPromoApp.__new__(GoldPromoApp)
        app.pending_etl = SimpleNamespace(
            path_src=(Path("a.xlsx"), Path("b.xlsm")),
            should_generate_so_sitegroup=False,
            _validate_exception_carton_discounts=MagicMock(),
            src=pd.DataFrame(
                {
                    "STRUCTURE": [110, 120],
                    "FILE NAME": ["a.xlsx", "b.xlsm"],
                    "_SOURCE_ROW": [8, 8],
                    "NOTE ERR FROM MASTER DATA": ["", ""],
                }
            ),
        )
        app.stage1_output = MagicMock()
        app.stage1_status = MagicMock()

        with TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory)
            with (
                patch.object(app, "_output_dir", return_value=output),
                patch.object(app, "_group_output_dir", return_value=output),
                patch.object(WorkbookExporter, "write_processed_source") as write_individual,
                patch.object(WorkbookExporter, "write_combined_processed_source") as write_combined,
                patch("src.desktop_app.messagebox.showinfo"),
            ):
                app.export_processed_src()

        self.assertEqual(write_individual.call_count, 2)
        write_combined.assert_called_once()
        self.assertEqual(write_combined.call_args.args[1].name[:14], "ALL_PROCESSED_")
        self.assertTrue(write_combined.call_args.kwargs["include_sitegroup"])

    def test_export_processed_source_does_not_add_combined_file_for_one_source(self):
        app = GoldPromoApp.__new__(GoldPromoApp)
        app.pending_etl = SimpleNamespace(
            path_src=(Path("a.xlsx"),),
            should_generate_so_sitegroup=True,
            _validate_exception_carton_discounts=MagicMock(),
            src=pd.DataFrame(
                {
                    "STRUCTURE": [110],
                    "FILE NAME": ["a.xlsx"],
                    "_SOURCE_ROW": [8],
                    "NOTE ERR FROM MASTER DATA": [""],
                }
            ),
        )
        app.stage1_output = MagicMock()
        app.stage1_status = MagicMock()

        with TemporaryDirectory(dir=Path.cwd()) as directory:
            output = Path(directory)
            with (
                patch.object(app, "_output_dir", return_value=output),
                patch.object(app, "_group_output_dir", return_value=output),
                patch.object(WorkbookExporter, "write_processed_source"),
                patch.object(WorkbookExporter, "write_combined_processed_source") as write_combined,
                patch("src.desktop_app.messagebox.showinfo"),
            ):
                app.export_processed_src()

        write_combined.assert_not_called()

    def test_export_processed_source_does_nothing_without_validated_batch(self):
        app = GoldPromoApp.__new__(GoldPromoApp)
        app.pending_etl = None
        app.stage1_output = MagicMock()

        with (
            patch.object(WorkbookExporter, "write_processed_source") as write_individual,
            patch.object(WorkbookExporter, "write_combined_processed_source") as write_combined,
        ):
            app.export_processed_src()

        write_individual.assert_not_called()
        write_combined.assert_not_called()


if __name__ == "__main__":
    unittest.main()
